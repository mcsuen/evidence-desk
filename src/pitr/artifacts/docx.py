"""Native Word text, tables, references and editable charts from ReportDocument.

Chart packaging follows ECMA-376 chart/externalData and the embedded workbook
relationship documented by Microsoft OfficeDev's ChartMarkup.xml sample.
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
import io
import zipfile

from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.part import Part, XmlPart
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
import xlsxwriter

from pitr.domain.common import digest
from pitr.domain.contracts import ValueSpan, CitationSpan, Paragraph, TableBlock
from pitr.domain.numbers import display

TEMPLATE = 'research.1'
COMPILER = 'docx.6'
FONT = 'Noto Sans CJK SC'
CHART_NS = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
DRAWING_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


def bookmark(paragraph, name, number):
    if len(name)>40:name=name[:15]+'_'+digest(name)[:24]
    paragraph._p.insert(0, element('w:bookmarkStart', **{'w:id': number, 'w:name': name}))
    paragraph._p.append(element('w:bookmarkEnd', **{'w:id': number}))


def field(paragraph, instruction, text='1'):
    begin, code, separate, value, end = [element('w:r') for _ in range(5)]
    begin.append(element('w:fldChar', **{'w:fldCharType': 'begin'}))
    it = element('w:instrText')
    it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    it.text = instruction
    code.append(it)
    separate.append(element('w:fldChar', **{'w:fldCharType': 'separate'}))
    t = element('w:t'); t.text = text; value.append(t)
    end.append(element('w:fldChar', **{'w:fldCharType': 'end'}))
    for item in (begin, code, separate, value, end):
        paragraph._p.append(item)


def hyperlink(paragraph, text, *, target=None, anchor=None):
    node = element('w:hyperlink')
    if target:
        node.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    if anchor:
        node.set(qn('w:anchor'), anchor)
    run = element('w:r')
    props = element('w:rPr'); props.append(element('w:color', **{'w:val': '315F81'})); run.append(props)
    t = element('w:t'); t.text = text; run.append(t); node.append(run)
    paragraph._p.append(node)


def native_chart(doc, chart, values, index, width):
    data = io.BytesIO()
    book = xlsxwriter.Workbook(data, {'in_memory': True})
    book.set_properties({'created': datetime(2000,1,1,tzinfo=timezone.utc), 'title': chart.title})
    sheet = book.add_worksheet('Data')
    sheet.write(0,0,'期间')
    for i, category in enumerate(chart.categories, 1):
        sheet.write_string(i,0,category)
    for j, series in enumerate(chart.series,1):
        sheet.write_string(0,j,series.name + ('（预测）' if series.role == 'forecast' else ''))
        for i, ref in enumerate(series.values,1):
            if ref:
                sheet.write_number(i,j,float(values[ref.key].amount))
    book.close()
    root = etree.Element('{'+CHART_NS+'}chartSpace', nsmap={'c':CHART_NS,'a':DRAWING_NS,'r':REL_NS})
    def c(parent, name, value=None, text=None):
        child = etree.SubElement(parent, '{'+CHART_NS+'}'+name)
        if value is not None: child.set('val',str(value))
        if text is not None: child.text = str(text)
        return child
    c(root, 'lang', 'zh-CN')
    body = c(root,'chart'); c(body,'autoTitleDeleted',1)
    plot = c(body,'plotArea'); c(plot,'layout')
    kind = c(plot,'barChart' if chart.kind == 'column' else 'lineChart')
    if chart.kind == 'column': c(kind,'barDir','col')
    c(kind,'grouping','clustered' if chart.kind == 'column' else 'standard')
    c(kind,'varyColors',0)
    colors = ['264B65','589BA2','B07D50','789568']
    points = []
    for j, series in enumerate(chart.series):
        ser = c(kind,'ser'); c(ser,'idx',j); c(ser,'order',j)
        tx = c(ser,'tx'); c(tx,'v',text=series.name+('（预测）' if series.role == 'forecast' else ''))
        style = c(ser,'spPr')
        fill = etree.SubElement(style,'{'+DRAWING_NS+'}solidFill')
        etree.SubElement(fill,'{'+DRAWING_NS+'}srgbClr',val=colors[j%len(colors)])
        line = etree.SubElement(style,'{'+DRAWING_NS+'}ln')
        color = etree.SubElement(line,'{'+DRAWING_NS+'}solidFill')
        etree.SubElement(color,'{'+DRAWING_NS+'}srgbClr',val=colors[j%len(colors)])
        if series.role == 'forecast': etree.SubElement(line,'{'+DRAWING_NS+'}prstDash',val='dash')
        cat = c(ser,'cat'); catref = c(cat,'strRef')
        c(catref,'f',text=f'Data!$A$2:$A${len(chart.categories)+1}')
        cache = c(catref,'strCache'); c(cache,'ptCount',len(chart.categories))
        for i, category in enumerate(chart.categories):
            pt = c(cache,'pt'); pt.set('idx',str(i)); c(pt,'v',text=category)
        val = c(ser,'val'); numref = c(val,'numRef')
        col = xlsxwriter.utility.xl_col_to_name(j+1)
        c(numref,'f',text=f'Data!${col}$2:${col}${len(chart.categories)+1}')
        cache = c(numref,'numCache'); c(cache,'formatCode',text='#,##0.000'); c(cache,'ptCount',len(series.values))
        for i, ref in enumerate(series.values):
            if ref:
                value = values[ref.key]
                pt = c(cache,'pt'); pt.set('idx',str(i)); c(pt,'v',text=str(value.amount))
                points.append({'series': j, 'index': i, 'ref': ref.model_dump(), 'amount': str(value.amount), 'display': display(value)})
        if chart.kind == 'line':
            marker = c(ser,'marker'); c(marker,'symbol','circle'); c(marker,'size',5)
            c(ser,'smooth',0)
    c(kind,'axId',1000+index*2); c(kind,'axId',1001+index*2)
    for tag, axis_id, cross, pos in [('catAx',1000+index*2,1001+index*2,'b'),('valAx',1001+index*2,1000+index*2,'l')]:
        axis = c(plot,tag); c(axis,'axId',axis_id)
        scaling = c(axis,'scaling'); c(scaling,'orientation','minMax')
        c(axis,'delete',0); c(axis,'axPos',pos)
        if tag == 'valAx':
            c(axis,'majorGridlines')
            fmt = c(axis,'numFmt'); fmt.set('formatCode','#,##0.##'); fmt.set('sourceLinked','0')
        c(axis,'majorTickMark','none'); c(axis,'minorTickMark','none'); c(axis,'tickLblPos','low' if tag=='catAx' else 'nextTo')
        c(axis,'crossAx',cross); c(axis,'crosses','autoZero')
        if tag == 'catAx': c(axis,'auto',1); c(axis,'lblAlgn','ctr'); c(axis,'lblOffset',100)
        else: c(axis,'crossBetween','between')
    legend = c(body,'legend'); c(legend,'legendPos','b'); c(legend,'overlay',0)
    c(body,'plotVisOnly',1); c(body,'dispBlanksAs','gap'); c(body,'showDLblsOverMax',0)
    # Font specified in chart text too, so CJK titles/categories do not fall back to missing glyphs.
    text = c(root,'txPr')
    etree.SubElement(text,'{'+DRAWING_NS+'}bodyPr'); etree.SubElement(text,'{'+DRAWING_NS+'}lstStyle')
    p = etree.SubElement(text,'{'+DRAWING_NS+'}p'); props=etree.SubElement(p,'{'+DRAWING_NS+'}pPr')
    defaults=etree.SubElement(props,'{'+DRAWING_NS+'}defRPr',sz='950')
    etree.SubElement(defaults,'{'+DRAWING_NS+'}latin',typeface=FONT)
    etree.SubElement(defaults,'{'+DRAWING_NS+'}ea',typeface=FONT)
    external = c(root,'externalData'); external.set('{'+REL_NS+'}id','rId1'); c(external,'autoUpdate',0)
    part = XmlPart(PackURI(f'/word/charts/chart{index}.xml'),
        'application/vnd.openxmlformats-officedocument.drawingml.chart+xml', root, doc.part.package)
    workbook = Part(PackURI(f'/word/embeddings/data{index}.xlsx'),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', data.getvalue(), doc.part.package)
    rid = part.relate_to(workbook, RT.PACKAGE)
    external.set('{'+REL_NS+'}id',rid)
    chart_rid = doc.part.relate_to(part,RT.CHART)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    drawing = parse_xml(f'''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
       xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
       xmlns:a="{DRAWING_NS}" xmlns:c="{CHART_NS}" xmlns:r="{REL_NS}">
       <wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{int(width)}" cy="2651760"/>
       <wp:docPr id="{index}" name="Chart {index}" descr="{chart.id}"/>
       <a:graphic><a:graphicData uri="{CHART_NS}"><c:chart r:id="{chart_rid}"/></a:graphicData></a:graphic>
       </wp:inline></w:drawing>''')
    paragraph.add_run()._r.append(drawing)
    return {'id': chart.id, 'kind': chart.kind, 'part': f'word/charts/chart{index}.xml', 'points': points,
            'workbook': f'word/embeddings/data{index}.xlsx', 'categories':chart.categories,
            'series':[s.name+('（预测）' if s.role=='forecast' else '') for s in chart.series],
            'workbook_digest': digest(data.getvalue()), 'unit': chart.unit}


def compile_docx(view, paper='Letter'):
    report = view.document
    values = {v.ref.key:v for v in view.values}
    evidence = {e.ref.key:e for e in view.evidence}
    sources = {s.ref.key:s for s in view.sources}
    references, counts, expected_text, charts = {}, Counter(), [], []
    bookmark_id = 0
    def mark(paragraph,name):
        nonlocal bookmark_id
        bookmark_id += 1
        bookmark(paragraph,name,bookmark_id)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = (Inches(8.5), Inches(11)) if paper == 'Letter' else (Mm(210),Mm(297))
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(18)
    section.header_distance = section.footer_distance = Mm(8)
    width = section.page_width-section.left_margin-section.right_margin
    for name, size in [('Normal',11),('Title',25),('Heading 1',19),('Heading 2',13),('Caption',9),('Subtitle',10)]:
        style = doc.styles[name]
        style.font.name = FONT; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string('18232D')
        props = style.element.get_or_add_rPr()
        fonts = props.rFonts
        if fonts is None: fonts = element('w:rFonts'); props.insert(0,fonts)
        for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key),FONT)
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith('theme'):del fonts.attrib[attribute]
        style.font.italic=False
        style.font.bold=name in ('Title','Heading 1','Heading 2')
        style.paragraph_format.space_after = Pt(7 if name == 'Normal' else 10)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.widow_control = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(18)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
    doc.core_properties.title = report.title
    doc.core_properties.author = 'PITR'
    doc.core_properties.subject = report.ref.key
    doc.core_properties.created = doc.core_properties.modified = datetime.fromisoformat(report.created_at)
    header = section.header.paragraphs[0]
    header.text = 'PITR  /  研究报告'
    header.style = doc.styles['Caption']
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f'修订 {report.revision}  ·  ')
    field(footer,' PAGE ')
    footer.add_run(' / '); field(footer,' NUMPAGES ')
    doc.add_heading(report.title,0)
    doc.add_paragraph(f'{report.created_at[:10]}  ·  修订 {report.revision}',style='Subtitle')
    # These status notes are frozen with the export, never recomputed inside a historical file.
    if view.delivery != 'ready':
        doc.add_paragraph('部分成果' if view.delivery == 'partial' else '待完成核验的草稿',style='Caption')
    def inlines(p, spans):
        text = ''
        for span in spans:
            if isinstance(span, ValueSpan):
                value = values[span.ref.key]; label = display(value)
                p.add_run(label); counts[span.ref.key] += 1; text += label
            elif isinstance(span, CitationSpan):
                if span.ref.key not in references: references[span.ref.key] = len(references)+1
                number = references[span.ref.key]
                label = f'[{number}]'; hyperlink(p,label,anchor='ref_'+str(number)); text += label
            else:
                p.add_run(span.text); text += span.text
        expected_text.append(text)
    def paragraph(block):
        p = doc.add_paragraph(); mark(p,'block_'+block.id)
        inlines(p,block.inlines)
    def table(block):
        doc.add_heading(block.title,2)
        tbl = doc.add_table(rows=1,cols=len(block.columns)); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        for column in tbl.columns: column.width = width//len(block.columns)
        tbl.rows[0]._tr.get_or_add_trPr().append(element('w:tblHeader'))
        props = tbl._tbl.tblPr
        borders = element('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            borders.append(element('w:'+side,**{'w:val':'single','w:sz':4,'w:color':'D8DFE4'}))
        props.append(borders)
        for i,label in enumerate(block.columns):
            cell=tbl.rows[0].cells[i]; cell.text=label
            cell._tc.get_or_add_tcPr().append(element('w:shd',**{'w:fill':'264B65'}))
            for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
        for row_index,row in enumerate(block.rows):
            cells=tbl.add_row().cells
            # Small rows stay together. Exceptionally long prose cells may span pages.
            if sum(len(getattr(s,'text','')) for cell in row for s in cell)<1200:
                cells[0]._tc.getparent().get_or_add_trPr().append(element('w:cantSplit'))
            for i,spans in enumerate(row):
                cell=cells[i];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index%2==0:cell._tc.get_or_add_tcPr().append(element('w:shd',**{'w:fill':'F2F5F7'}))
                inlines(cell.paragraphs[0],spans)
                if spans and all(isinstance(s,(ValueSpan,CitationSpan)) for s in spans):
                    cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph(block.note or '',style='Caption')
    doc.add_heading('研究结论',1)
    for p in report.summary: paragraph(p)
    for sec in report.sections:
        heading=doc.add_heading(sec.title,1);mark(heading,'section_'+sec.id)
        for block in sec.blocks:
            if isinstance(block,Paragraph):paragraph(block)
            elif isinstance(block,TableBlock):table(block)
            else:
                doc.add_heading(block.title,2)
                charts.append(native_chart(doc,block,values,len(charts)+1,width))
                doc.add_paragraph('单位：'+block.unit+('。'+block.note if block.note else ''),style='Caption')
    for title,items in [('研究限制与缺项',report.gaps),('下一步验证',report.next_steps)]:
        if items:
            doc.add_heading(title,1)
            for item in items:doc.add_paragraph(item,style='List Bullet')
    if references:
        doc.add_heading('原件与引用',1)
        for key,number in references.items():
            e=evidence[key];s=sources[e.source.key]
            p=doc.add_paragraph();mark(p,'ref_'+str(number))
            p.paragraph_format.keep_with_next=True
            p.add_run(f'[{number}] {s.title} · '+(f'第 {e.page} 页' if e.page else e.block_id)+f' · 原件修订 {s.revision}')
            quote=doc.add_paragraph(e.quote,style='Caption')
            if s.url.startswith(('https://','http://')):hyperlink(doc.add_paragraph(),s.url,target=s.url)
    # Fix zip metadata so retries from the same frozen content produce the same bytes.
    raw=io.BytesIO();doc.save(raw)
    normalized=io.BytesIO()
    with zipfile.ZipFile(raw) as source,zipfile.ZipFile(normalized,'w',compression=zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info=zipfile.ZipInfo(name,date_time=(2000,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED
            target.writestr(info,source.read(name))
    result=normalized.getvalue()
    manifest={'report':report.ref.model_dump(),'report_digest':digest(report),'paper':paper,
        'template':TEMPLATE,'compiler':COMPILER,'font':FONT,'value_occurrences':dict(counts),
        'citations':references,'charts':charts,'text':expected_text,'delivery_at_export':view.delivery,
        'checks_at_export':[c.model_dump(mode='json') for c in view.checks]}
    validate(result,manifest)
    return result,manifest


def validate(raw, manifest):
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','c':CHART_NS}
    with zipfile.ZipFile(io.BytesIO(raw)) as package:
        body=etree.fromstring(package.read('word/document.xml'))
        # Word text only; field instructions are not prose.
        paragraphs=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in body.xpath('.//w:p',namespaces=ns)]
        actual=Counter(paragraphs)
        expected=Counter(manifest['text'])
        if any(actual[text]<count for text,count in expected.items()):
            raise ValueError('DOCX 正文或表格与统一报告内容不一致')
        all_names=body.xpath('.//w:bookmarkStart/@w:name',namespaces=ns)
        ids=body.xpath('.//w:bookmarkStart/@w:id',namespaces=ns)
        if len(all_names)!=len(set(all_names)) or len(ids)!=len(set(ids)) or any(len(name)>40 for name in all_names):
            raise ValueError('DOCX 书签名称或标识无效')
        names=set(all_names)
        anchors=body.xpath('.//w:hyperlink/@w:anchor',namespaces=ns)
        if any(anchor not in names for anchor in anchors):
            raise ValueError('DOCX 存在断开的内部引用')
        for chart in manifest['charts']:
            tree=etree.fromstring(package.read(chart['part']))
            workbook=package.read(chart['workbook'])
            if digest(workbook)!=chart['workbook_digest']:raise ValueError('图表工作簿摘要不一致')
            with zipfile.ZipFile(io.BytesIO(workbook)) as book:
                xns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                cells=etree.fromstring(book.read('xl/worksheets/sheet1.xml'))
                strings=[''.join(si.itertext()) for si in etree.fromstring(book.read('xl/sharedStrings.xml'))]
                for i,label in enumerate(chart['categories'],2):
                    found=cells.xpath(f'//s:c[@r="A{i}"]/s:v/text()',namespaces=xns)
                    if not found or strings[int(found[0])]!=label:raise ValueError('工作簿期间与报告不一致')
                for j,label in enumerate(chart['series'],1):
                    address=xlsxwriter.utility.xl_rowcol_to_cell(0,j)
                    found=cells.xpath(f'//s:c[@r="{address}"]/s:v/text()',namespaces=xns)
                    if not found or strings[int(found[0])]!=label:raise ValueError('工作簿序列与报告不一致')
                expected_cells=set()
                for point in chart['points']:
                    address=xlsxwriter.utility.xl_rowcol_to_cell(point['index']+1,point['series']+1)
                    expected_cells.add(address)
                    found=cells.xpath(f'//s:c[@r="{address}"]/s:v/text()',namespaces=xns)
                    amount=Decimal(point['amount'])
                    # Spreadsheet numeric cells use IEEE-754; require agreement to 15 significant digits.
                    if not found or abs(Decimal(found[0])-amount)>max(abs(amount),Decimal(1))*Decimal('1e-15'):
                        raise ValueError('工作簿数值与报告不一致')
                numeric_cells={c.get('r') for c in cells.xpath('//s:c[not(@t)]',namespaces=xns)}
                if numeric_cells!=expected_cells:raise ValueError('工作簿缺项被填成了数值')
            for point in chart['points']:
                found=tree.xpath(f'//c:ser[c:idx/@val="{point["series"]}"]/c:val/c:numRef/c:numCache/c:pt[@idx="{point["index"]}"]/c:v/text()',namespaces=ns)
                if found != [point['amount']]:raise ValueError('图表数值与统一报告内容不一致')
    return {'status':'passed','coverage':['text','tables','citation_bookmarks','chart_values','native_workbooks']}
